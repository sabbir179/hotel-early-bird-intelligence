from pathlib import Path
from datetime import datetime
from typing import Dict, Optional, Union

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def format_currency(v) -> str:
    """Format value as currency."""
    try:
        return f"£{float(v):,.2f}"
    except Exception:
        return "-"


def format_percent(v, decimals: int = 1) -> str:
    """Format value as percentage."""
    try:
        return f"{float(v):.{decimals}f}%"
    except Exception:
        return "-"


def generate_powerpoint_report(
    df: pd.DataFrame,
    validation_counts: Optional[Dict[str, int]] = None,
    output_path: Optional[Union[str, Path]] = None
) -> Path:
    """Generate a PowerPoint report from dashboard data.
    
    Args:
        df: DataFrame with KPI data (must contain business_date and other metrics)
        validation_counts: Dict with validation status counts (optional)
        output_path: Custom output path (optional, defaults to reports/powerpoint/)
    
    Returns:
        Path to the generated PowerPoint file
    """
    if df is None or df.empty:
        raise ValueError("DataFrame is empty or None")
    
    # Get latest row
    latest = df.iloc[-1]
    latest_date = latest.get("business_date")
    
    if pd.isna(latest_date):
        latest_date_str = "Unknown"
    else:
        if not isinstance(latest_date, str):
            latest_date_str = pd.Timestamp(latest_date).strftime("%Y-%m-%d")
        else:
            latest_date_str = latest_date
    
    # Determine output path
    if output_path is None:
        report_dir = Path("reports") / "powerpoint"
        report_dir.mkdir(parents=True, exist_ok=True)
        filename = f"early_bird_report_{latest_date_str}.pptx"
        output_path = report_dir / filename
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    _add_title_slide(prs, latest_date_str)
    
    # Slide 2: Latest KPI Summary
    _add_kpi_summary_slide(prs, latest)
    
    # Slide 3: Management Summary
    _add_management_summary_slide(prs, df, latest)
    
    # Slide 4: Data Validation
    _add_validation_slide(prs, validation_counts)
    
    # Save
    prs.save(str(output_path))
    
    return output_path


def _add_title_slide(prs: Presentation, latest_date_str: str) -> None:
    """Add title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0, 51, 102)  # Dark blue
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = "Hotel Early Bird Intelligence"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    
    p = subtitle_frame.paragraphs[0]
    p.text = f"Latest report: {latest_date_str}"
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    p = subtitle_frame.add_paragraph()
    p.text = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER


def _add_kpi_summary_slide(prs: Presentation, latest_row) -> None:
    """Add KPI summary slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Latest KPI Summary"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    # KPI data
    kpi_data = [
        ("Business Date", latest_row.get("business_date") if not pd.isna(latest_row.get("business_date")) else "-"),
        ("Rooms Sold", int(latest_row.get("daily_rooms_sold")) if pd.notna(latest_row.get("daily_rooms_sold")) else "-"),
        ("Occupancy %", format_percent(latest_row.get("daily_occupancy_pct")) if pd.notna(latest_row.get("daily_occupancy_pct")) else "-"),
        ("ADR", format_currency(latest_row.get("daily_average_rate")) if pd.notna(latest_row.get("daily_average_rate")) else "-"),
        ("Total Revenue", format_currency(latest_row.get("revenue_mtd_total_revenue")) if pd.notna(latest_row.get("revenue_mtd_total_revenue")) else "-"),
        ("Rooms Revenue", format_currency(latest_row.get("revenue_stats_rooms_revenue")) if pd.notna(latest_row.get("revenue_stats_rooms_revenue")) else "-"),
        ("Bar Revenue", format_currency(latest_row.get("revenue_mtd_bar")) if pd.notna(latest_row.get("revenue_mtd_bar")) else "-"),
    ]
    
    # Create table
    rows, cols = len(kpi_data) + 1, 2
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(5)
    
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height).table
    
    # Set column widths
    table_shape.columns[0].width = Inches(4)
    table_shape.columns[1].width = Inches(4)
    
    # Header row
    header_cells = [table_shape.cell(0, 0), table_shape.cell(0, 1)]
    for cell in header_cells:
        cell.text_frame.clear()
        p = cell.text_frame.paragraphs[0]
        p.text = "Metric" if header_cells.index(cell) == 0 else "Value"
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.size = Pt(14)
        fill = cell.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 51, 102)
    
    # Data rows
    for idx, (metric, value) in enumerate(kpi_data, 1):
        cell = table_shape.cell(idx, 0)
        cell.text_frame.clear()
        p = cell.text_frame.paragraphs[0]
        p.text = metric
        p.font.size = Pt(12)
        
        cell = table_shape.cell(idx, 1)
        cell.text_frame.clear()
        p = cell.text_frame.paragraphs[0]
        p.text = str(value)
        p.font.size = Pt(12)
        p.alignment = PP_ALIGN.RIGHT


def _add_management_summary_slide(prs: Presentation, df: pd.DataFrame, latest_row) -> None:
    """Add management summary slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Management Summary"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    # Build summary items
    summary_items = []
    
    # Latest values
    latest_rooms = int(latest_row.get("daily_rooms_sold")) if pd.notna(latest_row.get("daily_rooms_sold")) else "-"
    summary_items.append(f"Latest rooms sold: {latest_rooms}")
    
    latest_occ = latest_row.get("daily_occupancy_pct") if "daily_occupancy_pct" in latest_row.index else None
    latest_occ_str = format_percent(latest_occ) if latest_occ is not None and pd.notna(latest_occ) else "-"
    summary_items.append(f"Latest occupancy: {latest_occ_str}")
    
    latest_adr = latest_row.get("daily_average_rate")
    latest_adr_str = format_currency(latest_adr) if pd.notna(latest_adr) else "-"
    summary_items.append(f"Latest ADR: {latest_adr_str}")
    
    latest_rev = latest_row.get("revenue_mtd_total_revenue")
    latest_rev_str = format_currency(latest_rev) if pd.notna(latest_rev) else "-"
    summary_items.append(f"Latest total revenue: {latest_rev_str}")
    
    # Compare with previous report if available
    if len(df) > 1:
        prev_date = df["business_date"].nlargest(2).iloc[-1]
        prev_row = df.loc[df["business_date"] == prev_date].tail(1)
        if not prev_row.empty:
            prev = prev_row.iloc[0]
            
            # Rooms sold change
            prev_rooms = prev.get("daily_rooms_sold")
            curr_rooms = latest_row.get("daily_rooms_sold")
            if pd.notna(prev_rooms) and pd.notna(curr_rooms):
                rooms_change = int(curr_rooms) - int(prev_rooms)
                rooms_change_str = f"+{rooms_change}" if rooms_change > 0 else str(rooms_change)
                summary_items.append(f"Rooms sold change: {rooms_change_str}")
            
            # Occupancy change
            prev_occ = prev.get("daily_occupancy_pct") if "daily_occupancy_pct" in prev.index else None
            if latest_occ is not None and pd.notna(latest_occ) and prev_occ is not None and pd.notna(prev_occ):
                occ_change = latest_occ - prev_occ
                occ_change_str = f"+{occ_change:.1f}" if occ_change > 0 else f"{occ_change:.1f}"
                summary_items.append(f"Occupancy change: {occ_change_str} pp")
            
            # ADR change
            prev_adr = prev.get("daily_average_rate")
            if pd.notna(prev_adr) and pd.notna(latest_adr):
                adr_change = latest_adr - prev_adr
                adr_change_str = f"+£{adr_change:.2f}" if adr_change > 0 else f"£{adr_change:.2f}"
                summary_items.append(f"ADR change: {adr_change_str}")
    else:
        summary_items.append("No previous report available for comparison")
    
    # Add summary items as bullet points
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5.5))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    for idx, item in enumerate(summary_items):
        if idx == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(18)
        p.level = 0
        p.space_before = Pt(8)


def _add_validation_slide(prs: Presentation, validation_counts: Optional[Dict[str, int]]) -> None:
    """Add validation slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Data Validation"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    # Validation summary
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(5))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    if validation_counts is None or len(validation_counts) == 0:
        p = text_frame.paragraphs[0]
        p.text = "No validation data available"
        p.font.size = Pt(18)
    else:
        p = text_frame.paragraphs[0]
        p.text = "Validation Status Counts:"
        p.font.size = Pt(20)
        p.font.bold = True
        p.space_after = Pt(12)
        
        for status, count in validation_counts.items():
            p = text_frame.add_paragraph()
            p.text = f"• {status}: {count}"
            p.font.size = Pt(18)
            p.level = 0
            p.space_before = Pt(6)
        
        # Summary line
        fail_count = validation_counts.get("FAIL", 0)
        warn_count = validation_counts.get("WARN", 0)
        
        p = text_frame.add_paragraph()
        p.text = ""
        p.space_before = Pt(12)
        
        if fail_count == 0 and warn_count == 0:
            p = text_frame.add_paragraph()
            p.text = "✓ All checks passed"
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 128, 0)
        else:
            p = text_frame.add_paragraph()
            p.text = f"⚠ {fail_count} failures, {warn_count} warnings"
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 165, 0)


class ReportGenerator:
    """Generate reports in PowerPoint and PDF formats."""

    def __init__(self, report_dir: str):
        self.report_dir = Path(report_dir)
        self.report_dir.mkdir(parents=True, exist_ok=True)

    def create_pptx(self, output_name: str) -> Path:
        """Create a PowerPoint report placeholder."""
        output_path = self.report_dir / output_name
        return output_path

    def create_pdf(self, output_name: str) -> Path:
        """Create a PDF report placeholder."""
        output_path = self.report_dir / output_name
        return output_path
