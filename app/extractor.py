import re
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from openpyxl import load_workbook

DATE_PATTERN = re.compile(r"(\d{2})\.(\d{2})\.(\d{2})")

TARGET_KPI_LABELS = [
    "Total Revenue (All)",
    "Rooms only revenue",
    "Rooms sold (physically filled overnight)",
    "ADR",
    "Out of Order Rooms",
    "Day Use Rooms /Early departure",
    "Total Occupancy including day use rooms",
    "No shows",
    "Book outs",
    "Complimentary rooms",
    "Total Food & Beverage Revenue",
    "GPK (F&B)",
    "Lobby Bar( F&B)",
    "M&E Room Hire Etc",
]

NORMALIZED_KPI_LABELS = {label: re.sub(r"\s+", " ", label.strip().lower()) for label in TARGET_KPI_LABELS}


def normalize_text(value: Any) -> str:
    """Normalize text for label matching."""
    if value is None:
        return ""

    text = str(value).strip().lower()
    return re.sub(r"\s+", " ", text)


def list_early_bird_files(raw_dir: Union[str, Path] = None) -> List[Path]:
    """Return sorted .xlsx files from the raw data folder, excluding temp files."""
    raw_path = Path(raw_dir) if raw_dir is not None else Path("data/raw")
    if not raw_path.exists() or not raw_path.is_dir():
        return []

    files = [
        path
        for path in raw_path.iterdir()
        if path.is_file()
        and path.suffix.lower() == ".xlsx"
        and not path.name.startswith("~$")
    ]
    return sorted(files, key=lambda path: path.name)


def extract_date_from_filename(file_path: Union[str, Path]) -> Optional[date]:
    """Extract a DD.MM.YY date from a filename and normalize the year."""
    file_name = Path(file_path).name
    match = DATE_PATTERN.search(file_name)
    if not match:
        return None

    day, month, year = match.groups()
    day_num = int(day)
    month_num = int(month)
    year_num = int(year)

    if 0 <= year_num <= 79:
        year_num += 2000
    else:
        year_num += 1900

    try:
        return date(year_num, month_num, day_num)
    except ValueError:
        return None


def find_value_to_right(ws, row: int, col: int, max_scan_columns: int = 10) -> Dict[str, Optional[Union[str, int, float]]]:
    """Scan to the right of a label cell for the first non-empty value."""
    for offset in range(1, max_scan_columns + 1):
        scan_col = col + offset
        if scan_col > ws.max_column:
            break

        candidate = ws.cell(row=row, column=scan_col)
        if candidate.value is None:
            continue

        normalized_value = normalize_text(candidate.value)
        if normalized_value == "":
            continue

        return {"value": candidate.value, "coordinate": candidate.coordinate}

    return {"value": None, "coordinate": None}


def find_section_start(ws, section_title: str):
    """Find the cell that exactly matches the normalized `section_title` in the worksheet.

    Returns the cell object when found, otherwise None.
    """
    if section_title is None:
        return None

    target = normalize_text(section_title)
    for row in ws.iter_rows():
        for cell in row:
            if normalize_text(cell.value) == target:
                return cell

    return None


def _coerce_to_number(value: Any) -> Optional[Union[int, float]]:
    """Try to coerce a cell value to a numeric type. Returns None if not numeric."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        # Remove common formatting like percent, commas, whitespace
        s = str(value).strip().replace("%", "").replace(",", "")
        if s == "":
            return None
        return float(s)
    except Exception:
        return None


def find_label_value_in_section(ws, section_title: str, label_text: str, max_rows: int = 15, max_scan_cols: int = 10) -> Optional[Union[int, float]]:
    """Find a numeric value for `label_text` located in the block headed by `section_title`.

    - Locates the section start cell by exact normalized match to `section_title`.
    - Scans downward up to `max_rows` rows in the same column looking for an exact
      normalized match to `label_text`.
    - For a matching label row, scans right up to `max_scan_cols` columns and
      returns the first numeric value found (int/float or coercible string).
    - Returns None if nothing numeric is found.
    """
    sec_cell = find_section_start(ws, section_title)
    if sec_cell is None:
        return None

    start_row = sec_cell.row
    col = sec_cell.column
    end_row = min(ws.max_row, start_row + max_rows)

    target_label = normalize_text(label_text)

    for r in range(start_row + 1, end_row + 1):
        cell = ws.cell(row=r, column=col)
        if normalize_text(cell.value) != target_label:
            continue

        # Found the label row; scan right for a numeric value
        for offset in range(1, max_scan_cols + 1):
            scan_col = col + offset
            if scan_col > ws.max_column:
                break
            candidate = ws.cell(row=r, column=scan_col)
            num = _coerce_to_number(candidate.value)
            if num is not None:
                return num

        # If label matched but no numeric to the right, stop searching further labels
        return None


def find_label_value_after_section(ws, section_title: str, label_text: str, max_rows: int = 20, max_scan_cols: int = 10) -> Optional[Union[int, float]]:
    """Find a numeric value for `label_text` located somewhere in rows after `section_title`.

    - Locates the section start cell by exact normalized match to `section_title`.
    - Scans each row downward up to `max_rows` to find a cell whose normalized
      text exactly matches `label_text` (label can be in any column).
    - From the matching label cell, scans right up to `max_scan_cols` columns
      and returns the first numeric value found.
    - Returns None if not found.
    """
    sec_cell = find_section_start(ws, section_title)
    if sec_cell is None:
        return None

    start_row = sec_cell.row
    end_row = min(ws.max_row, start_row + max_rows)
    target_label = normalize_text(label_text)

    for r in range(start_row + 1, end_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if normalize_text(cell.value) != target_label:
                continue

            # Found the label cell; scan right for numeric value
            for offset in range(1, max_scan_cols + 1):
                scan_col = c + offset
                if scan_col > ws.max_column:
                    break
                candidate = ws.cell(row=r, column=scan_col)
                num = _coerce_to_number(candidate.value)
                if num is not None:
                    return num

            return None




def inspect_workbook(file_path: Union[str, Path], max_rows: int = 30, max_cols: int = 15) -> None:
    """Debug helper: print workbook sheet names and preview non-empty cells.

    Prints exact sheet names and, for each worksheet, its dimensions and
    non-empty cells within the provided preview window.
    """
    wb = load_workbook(file_path, data_only=True)
    print("Sheets:", wb.sheetnames)

    for sheet in wb.worksheets:
        print(f"\nSheet: {sheet.title}")
        print(f"  max_row={sheet.max_row}, max_column={sheet.max_column}")
        rows_to_scan = min(max_rows, sheet.max_row)
        cols_to_scan = min(max_cols, sheet.max_column)

        for r in range(1, rows_to_scan + 1):
            for c in range(1, cols_to_scan + 1):
                cell = sheet.cell(row=r, column=c)
                if cell.value is not None and normalize_text(cell.value) != "":
                    print(f"    {cell.coordinate} | {cell.value}")

def get_cell_value(ws, cell_ref: str):
    """Return the raw value of the cell identified by `cell_ref`.

    Keeps numeric values (including percentage decimals) as-is. Returns None
    if the cell is empty or does not exist.
    """
    try:
        cell = ws[cell_ref]
    except Exception:
        return None

    return cell.value


def find_text_value_to_right(ws, label_text: str, max_scan_columns: int = 10):
    """Find text value to the right of a label cell.

    Locates a label by exact normalized text anywhere in the worksheet and
    returns the first non-empty value found to the right.
    """
    target = normalize_text(label_text)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if normalize_text(cell.value) != target:
                continue

            for offset in range(1, max_scan_columns + 1):
                scan_col = cell.column + offset
                if scan_col > ws.max_column:
                    break
                candidate = ws.cell(row=cell.row, column=scan_col)
                if candidate.value is None:
                    continue
                if normalize_text(candidate.value) == "":
                    continue
                return candidate.value
            return None
    return None


def find_text_value_near_label(ws, label_text: str, max_scan_columns: int = 10, max_scan_rows: int = 5):
    """Find a text value near a label cell, scanning right then below."""
    target = normalize_text(label_text)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if normalize_text(cell.value) != target:
                continue

            # Try right first
            for offset in range(1, max_scan_columns + 1):
                scan_col = cell.column + offset
                if scan_col > ws.max_column:
                    break
                candidate = ws.cell(row=cell.row, column=scan_col)
                if candidate.value is None:
                    continue
                if normalize_text(candidate.value) == "":
                    continue
                return candidate.value

            # Try below in the same column
            for offset in range(1, max_scan_rows + 1):
                scan_row = cell.row + offset
                if scan_row > ws.max_row:
                    break
                candidate = ws.cell(row=scan_row, column=cell.column)
                if candidate.value is None:
                    continue
                if normalize_text(candidate.value) == "":
                    continue
                return candidate.value

            return None
    return None


def extract_report_kpis(file_path: Union[str, Path]) -> Dict[str, Optional[object]]:
    """Extract a clean KPI dictionary from the Report sheet for a file.

    Cell mappings follow the inspected template. Empty cells are left as None.
    Raises ValueError if the Report sheet is missing.
    """
    wb = load_workbook(file_path, data_only=True)
    if "Report" not in wb.sheetnames:
        raise ValueError(f"Report sheet not found in {file_path}")

    ws = wb["Report"]
    result: Dict[str, Optional[object]] = {}
    result["file_name"] = Path(file_path).name
    result["business_date"] = extract_date_from_filename(file_path)

    # Daily
    result["daily_rooms_sold"] = get_cell_value(ws, "B13")
    result["daily_occupancy"] = get_cell_value(ws, "B14")
    result["daily_average_rate"] = get_cell_value(ws, "B15")
    result["house_comp_rooms"] = get_cell_value(ws, "B16")

    # MTD (and Business-on-books / BOB mapping as requested)
    result["mtd_rooms_sold"] = get_cell_value(ws, "G18")
    result["mtd_occupancy"] = get_cell_value(ws, "G19")
    result["mtd_rooms_revenue"] = get_cell_value(ws, "G20")
    result["mtd_arr"] = get_cell_value(ws, "G21")

    # BOB (business on books) - using same cells per template
    # BOB (business on books) - prefer locating values by section labels
    section_title = "Rooms Business on Books"
    rooms_bob_val = find_label_value_in_section(ws, section_title, "Rooms Sold")
    bob_occ_val = find_label_value_in_section(ws, section_title, "Occupancy")
    bob_rooms_rev_val = find_label_value_in_section(ws, section_title, "Rooms Revenue")
    bob_arr_val = find_label_value_in_section(ws, section_title, "ARR")
    bob_rest_val = find_label_value_in_section(ws, section_title, "BOB rest of month")

    result["rooms_business_on_books"] = rooms_bob_val if rooms_bob_val is not None else get_cell_value(ws, "G18")
    result["bob_occupancy"] = bob_occ_val if bob_occ_val is not None else get_cell_value(ws, "G19")
    result["bob_rooms_revenue"] = bob_rooms_rev_val if bob_rooms_rev_val is not None else get_cell_value(ws, "G20")
    result["bob_arr"] = bob_arr_val if bob_arr_val is not None else get_cell_value(ws, "G21")
    result["bob_rest_of_month"] = bob_rest_val if bob_rest_val is not None else get_cell_value(ws, "G22")

    # Revenue MTD block (K column)
    result["revenue_mtd_rooms_revenue"] = get_cell_value(ws, "K13")
    result["revenue_mtd_other_rooms_rev"] = get_cell_value(ws, "K14")
    result["revenue_mtd_tel_internet_movies"] = get_cell_value(ws, "K15")
    result["revenue_mtd_greenport"] = get_cell_value(ws, "K16")
    result["revenue_mtd_bar"] = get_cell_value(ws, "K17")
    result["revenue_mtd_other_fnb"] = get_cell_value(ws, "K18")
    result["revenue_mtd_me"] = get_cell_value(ws, "K19")
    result["revenue_mtd_total_revenue"] = get_cell_value(ws, "K21")
    result["revenue_mtd_bob_rooms"] = get_cell_value(ws, "K22")

    # Revenue stats block (B column)
    # Revenue stats block (B column) - prefer locating values by section labels
    rev_stats_section = "REVENUE & STATS"
    rev_rooms = find_label_value_after_section(ws, rev_stats_section, "Rooms Revenue")
    rev_other_rooms = find_label_value_after_section(ws, rev_stats_section, "Other Rooms Rev")
    rev_tel = find_label_value_after_section(ws, rev_stats_section, "Tel/Internet/Movies")

    result["revenue_stats_rooms_revenue"] = rev_rooms if rev_rooms is not None else get_cell_value(ws, "B27")
    result["revenue_stats_other_rooms_rev"] = rev_other_rooms if rev_other_rooms is not None else get_cell_value(ws, "B28")
    result["revenue_stats_tel_internet_movies"] = rev_tel if rev_tel is not None else get_cell_value(ws, "B30")

    # Duty Manager / DM names, if present in the report
    result["early_dm_name"] = None
    result["late_dm_name"] = None
    result["night_dm_name"] = None

    for field_label, key in [
        ("Early DM", "early_dm_name"),
        ("Late DM", "late_dm_name"),
        ("Night DM", "night_dm_name"),
    ]:
        dm_value = find_text_value_near_label(ws, field_label)
        result[key] = str(dm_value).strip() if dm_value is not None else None

    # Backward-compatible general DM field if a single duty manager label exists
    dm_name = result["early_dm_name"] or result["late_dm_name"] or result["night_dm_name"]
    if dm_name is None:
        dm_name = find_text_value_near_label(ws, "Duty Manager")
    if dm_name is None:
        dm_name = find_text_value_near_label(ws, "DM")
    if dm_name is None:
        dm_name = find_text_value_near_label(ws, "DM Name")
    result["dm_name"] = str(dm_name).strip() if dm_name is not None else None

    return result


def extract_all_reports(raw_dir: Union[str, Path] = None) -> List[Dict[str, Optional[object]]]:
    """Extract KPIs from all Early Bird files in `raw_dir` and return list of dicts.

    Files are sorted by filename discovery order and results sorted by `business_date`.
    """
    files = list_early_bird_files(raw_dir)
    records: List[Dict[str, Optional[object]]] = []
    for f in files:
        try:
            rec = extract_report_kpis(f)
            records.append(rec)
        except Exception as e:
            print(f"Warning: failed to extract {f.name}: {e}")
            continue

    # Sort by business_date when available
    records = sorted(records, key=lambda r: (r.get("business_date") or date.min))
    return records


def run_pipeline(raw_dir: Union[str, Path] = None) -> Dict[str, object]:
    """Run the full extraction -> CSV -> validation -> DB pipeline.

    Returns a dict with summary information and paths.
    """
    records = extract_all_reports(raw_dir)

    # Import here to avoid top-level circular imports when module is imported by dashboard
    from app.analytics import records_to_dataframe, save_master_csv, save_validation_csv
    from app.validation import validate_records
    from app.database import save_daily_kpis

    df = records_to_dataframe(records)
    master_path = save_master_csv(df)

    validation_df = validate_records(df)
    validation_path = save_validation_csv(validation_df)

    db_path = save_daily_kpis(df)

    status_counts = validation_df["status"].value_counts().to_dict() if not validation_df.empty else {}

    return {
        "records_processed": len(records),
        "master_csv_path": master_path,
        "validation_csv_path": validation_path,
        "database_path": db_path,
        "validation_status_counts": status_counts,
    }

def find_label_cells(file_path: Union[str, Path]) -> List[Dict[str, Any]]:
    """Search workbook worksheets for target KPI labels."""
    workbook = load_workbook(file_path, data_only=True)
    matches: List[Dict[str, Any]] = []

    # Focus only on the main report worksheet to avoid noisy sheets
    if "Report" not in workbook.sheetnames:
        return []

    sheet = workbook["Report"]

    for row in sheet.iter_rows():
        for cell in row:
            normalized_text = normalize_text(cell.value)
            if not normalized_text:
                continue

            # Match when the normalized text equals the target label,
            # or the target label is contained inside the cell text.
            for label in TARGET_KPI_LABELS:
                normalized_label = re.sub(r"\s+", " ", label.strip().lower())
                if normalized_text == normalized_label or normalized_label in normalized_text:
                    value_info = find_value_to_right(sheet, cell.row, cell.column)
                    matches.append(
                        {
                            "label": label,
                            "sheet": sheet.title,
                            "label_coordinate": cell.coordinate,
                            "value": value_info["value"],
                            "value_coordinate": value_info["coordinate"],
                        }
                    )
                    break

    return matches


class DataExtractor:
    """Placeholder for Excel extraction logic."""

    def __init__(self, raw_data_dir: str):
        self.raw_data_dir = Path(raw_data_dir)

    def discover_files(self) -> List[Path]:
        """Discover daily Excel files in the raw data directory."""
        return list_early_bird_files(self.raw_data_dir)

    def extract(self) -> List[dict]:
        """Extract records from discovered files."""
        return []


if __name__ == "__main__":
    # Toggle debug inspection mode for workbook diagnosis
    DEBUG_INSPECT = False

    files = list_early_bird_files()
    if not files:
        print("No .xlsx files found in data/raw/")
    else:
        if DEBUG_INSPECT:
            # Inspect only the first file for debugging and stop
            first = files[0]
            header = "=" * 27
            print(header)
            print(f"Inspecting {first.name} -> {extract_date_from_filename(first)}")
            print(header)
            inspect_workbook(first)
        else:
                # Aggregate all records and run pipeline
                try:
                    summary = run_pipeline()
                    print(f"Processed {summary['records_processed']} records")
                    print(f"Master CSV: {summary['master_csv_path']}")
                    print(f"Validation CSV: {summary['validation_csv_path']}")
                    print(f"SQLite DB: {summary['database_path']}")
                    print(f"Status counts: {summary['validation_status_counts']}")
                except Exception as e:
                    print(f"Pipeline failed: {e}")
